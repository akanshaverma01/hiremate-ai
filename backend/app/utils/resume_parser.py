"""
Extracts raw text from an uploaded resume file (PDF or DOCX).

Known limitation: layouts with multi-column formatting, text boxes, or
scanned/image-only PDFs will extract poorly or return empty text. This is a
real constraint of pdfplumber/python-docx, not something the ML step can fix.
"""
import os
import pdfplumber
from docx import Document


def extract_text_from_pdf(file_path: str) -> str:
    text_chunks = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_path: str) -> str:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # also pull text from tables, since resumes often use them for layout
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def extract_resume_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not text or not text.strip():
        raise ValueError(
            "Could not extract text from this file. It may be a scanned/image-based "
            "document or use an unsupported layout."
        )
    return text
